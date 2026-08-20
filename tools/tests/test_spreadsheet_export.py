import sys
import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from spreadsheet_export import write_leads


class SpreadsheetExportTests(unittest.TestCase):
    def test_writes_formatted_workbook(self):
        rows = [{
            "firm_name": "Wasatch Legal",
            "website": "https://wasatch.test/",
            "business_emails": "info@wasatch.test",
            "phones": "(801) 555-1234",
            "estimated_lawyers": "4",
            "firm_size": "small (1-10 lawyers)",
        }]
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "leads.xlsx"
            write_leads(rows, output, list(rows[0]), "family law firms in Provo, Utah")
            workbook = load_workbook(output, data_only=False)
            self.assertEqual(workbook.sheetnames, ["Overview", "Leads"])
            self.assertEqual(workbook["Leads"]["A2"].value, "Wasatch Legal")
            self.assertEqual(workbook["Leads"]["B2"].hyperlink.target, "https://wasatch.test/")
            self.assertEqual(
                workbook["Overview"]["B6"].value,
                "Filter the Leads tab by Firm Size, Practice Areas, or Count Confidence.",
            )
            self.assertFalse(
                any(
                    isinstance(cell.value, str) and cell.value.startswith("=")
                    for sheet in workbook.worksheets
                    for row in sheet.iter_rows()
                    for cell in row
                )
            )
            self.assertEqual(workbook["Leads"].freeze_panes, "A2")
            self.assertIn("LawFirmLeads", workbook["Leads"].tables)

    def test_rejects_unknown_extension(self):
        with tempfile.TemporaryDirectory() as folder:
            with self.assertRaisesRegex(ValueError, "xlsx or .csv"):
                write_leads([], Path(folder) / "leads.txt", ["firm_name"])

    def test_formula_like_values_are_literal_in_xlsx_and_csv(self):
        dangerous = ["=1+1", "+SUM(A1:A2)", "-2+3", "@SUM(A1:A2)", "\t=1+1", "\r=1+1"]
        rows = [{"firm_name": value, "website": ""} for value in dangerous]
        with tempfile.TemporaryDirectory() as folder:
            xlsx_output = Path(folder) / "leads.xlsx"
            csv_output = Path(folder) / "leads.csv"
            write_leads(rows, xlsx_output, ["firm_name", "website"], "=malicious-query")
            write_leads(rows, csv_output, ["firm_name", "website"])

            workbook = load_workbook(xlsx_output, data_only=False)
            cells = [workbook["Overview"]["B3"]] + [
                workbook["Leads"].cell(row=index, column=1)
                for index in range(2, len(dangerous) + 2)
            ]
            self.assertTrue(all(cell.data_type != "f" for cell in cells))
            self.assertTrue(all(str(cell.value).startswith("'") for cell in cells))

            with csv_output.open(newline="", encoding="utf-8-sig") as handle:
                saved = list(csv.DictReader(handle))
            self.assertTrue(all(row["firm_name"].startswith("'") for row in saved))


if __name__ == "__main__":
    unittest.main()
