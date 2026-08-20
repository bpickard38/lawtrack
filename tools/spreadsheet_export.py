"""Export researched law-firm leads to CSV or a formatted Excel workbook."""

from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


HEADER_FILL = PatternFill("solid", fgColor="17365D")
ACCENT_FILL = PatternFill("solid", fgColor="D9EAF7")
WHITE_BOLD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10, color="000000")
TITLE_FONT = Font(name="Arial", size=18, bold=True, color="17365D")
LINK_FONT = Font(name="Arial", size=10, color="0563C1", underline="single")


DISPLAY_HEADERS = {
    "firm_name": "Firm Name",
    "website": "Website",
    "business_emails": "Business Emails",
    "other_public_emails_review": "Other Public Emails — Review",
    "phones": "Phone Numbers",
    "estimated_lawyers": "Estimated Lawyers",
    "firm_size": "Firm Size",
    "lawyer_count_confidence": "Count Confidence",
    "practice_areas": "Practice Areas",
    "pages_crawled": "Pages Checked",
    "contact_source_urls": "Contact Sources",
    "lawyer_count_source_urls": "Lawyer Count Sources",
    "verified_date": "Verified Date",
    "notes": "Notes",
}


def spreadsheet_safe(value: object) -> object:
    """Force untrusted formula-like spreadsheet text to remain literal."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def write_csv(rows: list[dict[str, str]], output: Path, fieldnames: list[str]) -> None:
    with output.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(
            {field: spreadsheet_safe(row.get(field, "")) for field in fieldnames}
            for row in rows
        )


def write_xlsx(
    rows: list[dict[str, str]], output: Path, fieldnames: list[str], search_query: str = ""
) -> None:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    leads = workbook.create_sheet("Leads")

    overview["A1"] = "LawTrack Prospect Research"
    overview["A1"].font = TITLE_FONT
    overview["A3"] = "Search"
    overview["B3"] = spreadsheet_safe(search_query or "Website list supplied by user")
    overview["A4"] = "Generated"
    overview["B4"] = datetime.now().astimezone().strftime("%Y-%m-%d %H:%M %Z")
    overview["A6"] = "How to use"
    overview["B6"] = "Filter the Leads tab by Firm Size, Practice Areas, or Count Confidence."
    overview["A7"] = "Email guidance"
    overview["B7"] = "Prioritize Business Emails. Manually review other published addresses."
    overview["A8"] = "Important"
    overview["B8"] = "Lawyer counts are website-derived estimates. Verify before outreach."
    for row in range(3, 9):
        overview[f"A{row}"].font = Font(name="Arial", size=10, bold=True)
        overview[f"B{row}"].font = BODY_FONT
    for row in range(6, 9):
        overview[f"A{row}"].fill = ACCENT_FILL
        overview[f"B{row}"].fill = ACCENT_FILL
    overview.column_dimensions["A"].width = 22
    overview.column_dimensions["B"].width = 68
    overview.freeze_panes = "A3"

    for column_index, fieldname in enumerate(fieldnames, start=1):
        cell = leads.cell(row=1, column=column_index, value=DISPLAY_HEADERS.get(fieldname, fieldname))
        cell.fill = HEADER_FILL
        cell.font = WHITE_BOLD
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row_index, row in enumerate(rows, start=2):
        for column_index, fieldname in enumerate(fieldnames, start=1):
            value: object = spreadsheet_safe(row.get(fieldname, ""))
            if fieldname in {"estimated_lawyers", "pages_crawled"} and str(value).isdigit():
                value = int(value)
            cell = leads.cell(row=row_index, column=column_index, value=value)
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if fieldname == "website" and isinstance(value, str) and value.startswith(("http://", "https://")):
                cell.hyperlink = value
                cell.font = LINK_FONT

    leads.freeze_panes = "A2"
    leads.auto_filter.ref = leads.dimensions
    leads.row_dimensions[1].height = 32
    widths = {
        "firm_name": 30,
        "website": 34,
        "business_emails": 28,
        "other_public_emails_review": 30,
        "phones": 22,
        "estimated_lawyers": 16,
        "firm_size": 24,
        "lawyer_count_confidence": 18,
        "practice_areas": 32,
        "pages_crawled": 14,
        "contact_source_urls": 45,
        "lawyer_count_source_urls": 45,
        "verified_date": 15,
        "notes": 48,
    }
    for column_index, fieldname in enumerate(fieldnames, start=1):
        leads.column_dimensions[leads.cell(1, column_index).column_letter].width = widths.get(fieldname, 24)

    if rows:
        table = Table(displayName="LawFirmLeads", ref=f"A1:{leads.cell(len(rows) + 1, len(fieldnames)).coordinate}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        leads.add_table(table)

    workbook.save(output)


def write_leads(
    rows: list[dict[str, str]], output: Path, fieldnames: list[str], search_query: str = ""
) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    suffix = output.suffix.lower()
    if suffix == ".xlsx":
        write_xlsx(rows, output, fieldnames, search_query)
    elif suffix == ".csv":
        write_csv(rows, output, fieldnames)
    else:
        raise ValueError("Output filename must end in .xlsx or .csv")
